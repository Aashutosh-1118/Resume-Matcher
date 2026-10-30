import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.core.config import settings

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class DocumentError(ValueError):
    """A problem with an uploaded file that the user can fix."""


def read_upload(upload) -> tuple[str, bytes]:
    """Read a FastAPI UploadFile with a hard size cap. Returns (filename, bytes)."""
    name = Path(upload.filename or "file").name
    data = upload.file.read(settings.max_upload_bytes + 1)
    if len(data) > settings.max_upload_bytes:
        raise DocumentError(f"{name}: file is larger than {settings.max_upload_mb} MB.")
    if not data:
        raise DocumentError(f"{name}: file is empty.")
    return name, data


def _read_pdf(data: bytes) -> str:
    if not data.lstrip().startswith(b"%PDF"):
        raise DocumentError("This file is not a valid PDF.")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise DocumentError("This PDF is password protected.")
        pages = reader.pages[: settings.max_pdf_pages]
        parts = []
        for page in pages:
            text = page.extract_text()
            if text:
                parts.append(text)
        return "\n".join(parts)
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("Could not read this PDF. It may be damaged.")


def _read_docx(data: bytes) -> str:
    if not data.startswith(b"PK"):
        raise DocumentError("This file is not a valid DOCX.")
    try:
        document = Document(io.BytesIO(data))
    except Exception:
        raise DocumentError("Could not read this DOCX. It may be damaged.")

    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            seen = []
            for cell in row.cells:
                text = cell.text.strip()
                # Merged cells repeat the same text, so skip consecutive duplicates.
                if text and (not seen or seen[-1] != text):
                    seen.append(text)
            lines.extend(seen)
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    """Turn an uploaded PDF, DOCX or TXT into plain text."""
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise DocumentError(f"{filename}: only PDF, DOCX and TXT files are supported.")

    if suffix == ".pdf":
        text = _read_pdf(data)
    elif suffix == ".docx":
        text = _read_docx(data)
    else:
        text = data.decode("utf-8", errors="ignore")

    text = text.strip()
    if not text:
        raise DocumentError(
            f"{filename}: no readable text found. Scanned or image-only files are not supported yet."
        )
    return text
